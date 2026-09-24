"""Carrega o modelo de contabilização padrão BHub (plano de contas + rubricas por tipo).

Reconstruído em 2026-08-19 a partir de projetos/rubricas-folha - Versão base/notas.md
(o código original foi perdido; esta versão foi validada apenas contra a estrutura real
das planilhas de referência em app/dados_bhub/, não contra uma rodada completa do
pipeline). Ver notas.md para o histórico de decisões e regras de negócio confirmadas
com o Robert.
"""

from dataclasses import dataclass, field, replace

import openpyxl

NATUREZAS_EMPRESA = [
    "cpv_fabrica",
    "csv_servicos",
    "despesa_vendas",
    "despesa_administrativa",
    "cmv",
]

# Rótulo de exibição de cada natureza — CPV/CSV/CMV são siglas (Custo do Produto
# Vendido, Custo do Serviço Vendido, Custo da Mercadoria Vendida) e não devem
# aparecer em minúsculo; um `.title()` genérico erra isso (viraria "Cpv", "Csv").
NATUREZAS_EMPRESA_LABELS = {
    "cpv_fabrica": "CPV (Fábrica)",
    "csv_servicos": "CSV (Serviços)",
    "despesa_vendas": "Despesa com Vendas",
    "despesa_administrativa": "Despesa Administrativa",
    "cmv": "CMV",
}

SHEETS_RUBRICAS = {
    "folha": "Folha",
    "rescisao": "Rescisão",
    "empresa": "Empresa",
    "provisao_ferias": "Provisão de Férias",
    "provisao_13": "Provisão de 13º",
}

ABA_PLANO_CONTAS = "Plano de contas Padrão"
ABA_RELACAO_ESOCIAL = "Relação Base 141.935"
CODI_EMP_BASE = "100000"

# Tipo da Integração no EVENTO/INTEGRA do Domínio -> tipo de rubrica no modelo BHub.
# Tipo 3 (Férias) não tem aba própria: confirmado com o Robert que a maioria das
# rubricas de férias é contabilizada como tipo 1 (Folha mensal) mesmo — por isso não
# entra neste mapa; a UI deve orientar a usar o tipo 1 por padrão.
TIPO_INTEGRACAO = {
    1: "folha",
    2: "empresa",
    4: "rescisao",
    5: "provisao_ferias",
    6: "provisao_13",
}

TIPO_INTEGRACAO_LABELS = {
    1: "1 - Folha mensal",
    2: "2 - Empresa",
    3: "3 - Férias (use o tipo 1, ver aviso)",
    4: "4 - Rescisão",
    5: "5 - Provisão de Férias",
    6: "6 - Provisão de 13º",
}

# Contrapartida de provisão de férias/13º (e encargos INSS/FGTS/PIS sobre elas) que,
# no tipo de integração 4 (Rescisão), deve ser substituída por "Rescisão a Pagar".
CONTAS_PROVISAO_PARA_RESCISAO = {"2029", "2030", "2031", "2032", "2033", "2034", "2035", "2036"}
CONTA_RESCISAO_A_PAGAR = "2006"


@dataclass
class ContaBhub:
    codigo: str
    nome: str
    grupo_da_conta: str
    natureza: str  # ativo / passivo / receita / despesa / patrimonio_liquido / desconhecida


@dataclass
class RubricaBhub:
    numero: str
    nome: str
    natureza: str  # P (provento) / D (desconto) / I / ID (informativo) / "-" (empresa/provisão)
    tipo: str  # folha / rescisao / empresa / provisao_ferias / provisao_13
    contas: dict  # natureza_empresa -> (conta_debito, conta_credito)
    codigo_esocial: str = ""


@dataclass
class ModeloBhub:
    plano_contas: dict  # codigo -> ContaBhub
    rubricas: list  # list[RubricaBhub]

    def contas_por_natureza(self, natureza):
        return [c for c in self.plano_contas.values() if c.natureza == natureza]

    def rubricas_por_tipo(self, tipo):
        return [r for r in self.rubricas if r.tipo == tipo]


def _limpar_codigo(valor):
    if valor is None:
        return ""
    texto = str(valor).strip()
    if texto.endswith(".0"):
        texto = texto[:-2]
    return texto


def derivar_natureza_por_nome(nome_grupo):
    """Classifica ativo/passivo/receita/despesa pelo NOME do grupo, não pelo dígito
    inicial do código — dígitos iniciais têm significados diferentes entre planos de
    contas de empresas diferentes (achado registrado em notas.md 2026-07-14c)."""
    nome_grupo = (nome_grupo or "").upper()
    if "ATIVO" in nome_grupo:
        return "ativo"
    if "PASSIVO" in nome_grupo:
        return "passivo"
    if "PATRIMÔNIO" in nome_grupo or "PATRIMONIO" in nome_grupo:
        return "patrimonio_liquido"
    if "RECEITA" in nome_grupo:
        return "receita"
    if "DESPESA" in nome_grupo or "CUSTO" in nome_grupo:
        return "despesa"
    return "desconhecida"


def _indice_colunas(ws, linha=1):
    indices = {}
    for celula in ws[linha]:
        if celula.value is not None:
            indices[str(celula.value).strip().lower()] = celula.column
    return indices


def carregar_plano_contas(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = _obter_aba(wb, ABA_PLANO_CONTAS)
    idx = _indice_colunas(ws, 1)
    col_codigo = idx["codigo_conta"]
    col_nome = idx["nome_conta"]
    col_tipo = idx["tipo_conta"]
    col_grupo = idx["grupo da conta"]

    plano = {}
    for linha in ws.iter_rows(min_row=2, values_only=False):
        tipo_conta = linha[col_tipo - 1].value
        if str(tipo_conta or "").strip().upper() != "A":
            continue  # só contas analíticas entram no matching
        codigo = _limpar_codigo(linha[col_codigo - 1].value)
        if not codigo:
            continue
        nome = str(linha[col_nome - 1].value or "").strip()
        grupo = str(linha[col_grupo - 1].value or "").strip()
        plano[codigo] = ContaBhub(
            codigo=codigo,
            nome=nome,
            grupo_da_conta=grupo,
            natureza=derivar_natureza_por_nome(grupo),
        )
    return plano


def _obter_aba(wb, nome):
    nome_norm = nome.strip().lower()
    for titulo in wb.sheetnames:
        if titulo.strip().lower() == nome_norm:
            return wb[titulo]
    raise KeyError(f"Aba '{nome}' não encontrada. Abas disponíveis: {wb.sheetnames}")


def _localizar_linha_cabecalho(ws, max_linhas=15):
    for i in range(1, max_linhas + 1):
        for celula in ws[i]:
            if isinstance(celula.value, str) and celula.value.strip().lower() == "nome da rubrica":
                return i
    raise ValueError(f"Cabeçalho 'Nome da Rubrica' não encontrado na aba '{ws.title}'")


def mapear_colunas_rubricas(ws, linha_cab):
    """Acha as colunas Nº/Nome da Rubrica/Natureza e os pares Débito/Crédito
    (um por natureza da empresa, na ORDEM em que aparecem na planilha — a
    correspondência com NATUREZAS_EMPRESA é posicional, não por rótulo) na
    linha de cabeçalho `linha_cab`. Extraído à parte (não só usado aqui) pra
    bhub_model_editor.py reaproveitar exatamente a mesma lógica ao editar a
    planilha célula a célula, em vez de duplicar/arriscar divergir dela."""
    max_col = ws.max_column

    col_numero = col_nome = col_natureza = None
    pares_natureza = []
    par_atual = None

    for col in range(1, max_col + 1):
        titulo = ws.cell(row=linha_cab, column=col).value
        titulo_norm = str(titulo).strip().lower() if titulo else ""
        if titulo_norm == "nº":
            col_numero = col
        elif titulo_norm == "nome da rubrica":
            col_nome = col
        elif titulo_norm == "natureza":
            col_natureza = col
        elif titulo_norm == "débito":
            par_atual = {"debito": col}
        elif titulo_norm == "crédito" and par_atual is not None:
            par_atual["credito"] = col
            pares_natureza.append(par_atual)
            par_atual = None

    if col_numero is None or col_nome is None:
        raise ValueError(f"Colunas Nº/Nome da Rubrica não encontradas na aba '{ws.title}'")

    return col_numero, col_nome, col_natureza, pares_natureza


def _carregar_aba_rubricas(ws, tipo):
    linha_cab = _localizar_linha_cabecalho(ws)
    col_numero, col_nome, col_natureza, pares_natureza = mapear_colunas_rubricas(ws, linha_cab)

    rubricas = []
    for linha in range(linha_cab + 1, ws.max_row + 1):
        numero_valor = ws.cell(row=linha, column=col_numero).value
        if numero_valor in (None, ""):
            continue
        nome = ws.cell(row=linha, column=col_nome).value or ""
        natureza = ws.cell(row=linha, column=col_natureza).value if col_natureza else "-"

        contas = {}
        for slot_nome, par in zip(NATUREZAS_EMPRESA, pares_natureza):
            deb = ws.cell(row=linha, column=par["debito"]).value
            cred = ws.cell(row=linha, column=par["credito"]).value
            deb, cred = _limpar_codigo(deb), _limpar_codigo(cred)
            if deb and cred:
                contas[slot_nome] = (deb, cred)

        rubricas.append(RubricaBhub(
            numero=_limpar_codigo(numero_valor),
            nome=str(nome).strip(),
            natureza=str(natureza or "-").strip(),
            tipo=tipo,
            contas=contas,
        ))
    return rubricas


def _carregar_esocial_por_numero(ws):
    idx = _indice_colunas(ws, 1)
    col_i_eventos = idx.get("i_eventos")
    col_codi_emp = idx.get("codi_emp")
    col_esocial = idx.get("codigo_esocial")
    if not (col_i_eventos and col_codi_emp and col_esocial):
        return {}
    resultado = {}
    for linha in ws.iter_rows(min_row=2, values_only=False):
        codi_emp = _limpar_codigo(linha[col_codi_emp - 1].value)
        if codi_emp != CODI_EMP_BASE:
            continue
        numero = _limpar_codigo(linha[col_i_eventos - 1].value)
        codigo_esocial = _limpar_codigo(linha[col_esocial - 1].value)
        if numero:
            resultado[numero] = codigo_esocial
    return resultado


def load_bhub_model(caminho_parametrizacao, caminho_plano_contas):
    plano_contas = carregar_plano_contas(caminho_plano_contas)

    # (sem read_only=True: esta pasta de trabalho tem desenhos/células mescladas e
    # o modo read_only do openpyxl às vezes não calcula max_row/max_column)
    wb = openpyxl.load_workbook(caminho_parametrizacao, data_only=True)
    rubricas = []
    for tipo, nome_aba in SHEETS_RUBRICAS.items():
        ws = _obter_aba(wb, nome_aba)
        rubricas.extend(_carregar_aba_rubricas(ws, tipo))

    # O código eSocial só é confiável como chave de equivalência para rubricas de
    # Folha mensal — é o recorte que a "Relação Base 141.935" documenta (ver notas.md).
    ws_relacao = _obter_aba(wb, ABA_RELACAO_ESOCIAL)
    esocial_por_numero = _carregar_esocial_por_numero(ws_relacao)
    for r in rubricas:
        if r.tipo == "folha" and r.numero in esocial_por_numero:
            r.codigo_esocial = esocial_por_numero[r.numero]

    return ModeloBhub(plano_contas=plano_contas, rubricas=rubricas)


def substituir_provisao_por_rescisao(rubricas):
    """Regra confirmada: no Tipo de Integração 4 (Rescisão), contas de apropriação de
    férias/13º e seus encargos (2029-2036) são substituídas por 2006 (Rescisão a
    Pagar), tanto no lado débito quanto crédito."""
    resultado = []
    for r in rubricas:
        contas_ajustadas = {}
        for natureza_emp, (debito, credito) in r.contas.items():
            debito2 = CONTA_RESCISAO_A_PAGAR if debito in CONTAS_PROVISAO_PARA_RESCISAO else debito
            credito2 = CONTA_RESCISAO_A_PAGAR if credito in CONTAS_PROVISAO_PARA_RESCISAO else credito
            contas_ajustadas[natureza_emp] = (debito2, credito2)
        resultado.append(replace(r, contas=contas_ajustadas))
    return resultado
